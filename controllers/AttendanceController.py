import io
import logging
import re
from datetime import date, datetime
from functools import wraps
from urllib.parse import quote

from flask import render_template, request, jsonify, session, redirect, url_for, send_file
from models.AttendanceModel import AttendanceModel, STATUSES
from models.StudentModel import StudentModel

logger = logging.getLogger(__name__)

STUDENT_NO_RE = re.compile(r'^[A-Za-z0-9]+$')
DUPLICATE_MSG = '该学生当天已有考勤记录，请到列表中编辑原记录'


def _is_duplicate_error(e: Exception) -> bool:
    """MySQL 唯一约束冲突的错误码是 1062。并发下预检查可能漏过，靠数据库约束兜底。"""
    return bool(getattr(e, 'args', None)) and e.args[0] == 1062


def parse_date(value: str):
    """把 YYYY-MM-DD（也接受 2026-9-5）解析成 date，不合法返回 None。"""
    try:
        return datetime.strptime((value or '').strip(), '%Y-%m-%d').date()
    except ValueError:
        return None


def login_required(api: bool = False):
    """页面未登录跳登录页；接口未登录返回 401 JSON。"""
    def deco(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            if not session.get('teacher_id'):
                if api:
                    return jsonify({'code': 401, 'msg': '未登录'})
                return redirect(url_for('web.admin_login_page'))
            return fn(*args, **kwargs)
        return wrapper
    return deco


def validate_entry(data):
    """校验单条考勤的 状态/时间/备注（批量点名和单条录入共用）。返回 (entry, errors)。"""
    def text(key):
        return (data.get(key) or '').strip()

    entry = {'status': text('status') or '正常', 'time': text('time'), 'remark': text('remark')}
    errors = []
    if entry['status'] not in STATUSES:
        errors.append('状态只能是：' + ' / '.join(STATUSES))
    if len(entry['time']) > 20:
        errors.append('时间最多 20 个字')
    if len(entry['remark']) > 200:
        errors.append('备注最多 200 个字')
    return entry, errors


def validate_record(data):
    """校验并清洗单条考勤表单。返回 (record, errors)：errors 非空时 record 不可用。

    前端表单的限制只是提示，绕过页面直接调接口也必须被挡住，所以以这里为准。
    姓名不再信任前端提交的值，一律从学生名册取，保证与名册一致。
    """
    student_no = (data.get('student_no') or '').strip()
    date_raw = (data.get('date') or '').strip()
    entry, errors = validate_entry(data)
    record = {'student_no': student_no, 'student_name': '', 'date_str': date_raw, **entry}

    if not student_no:
        errors.insert(0, '请选择学生')
    elif len(student_no) > 20 or not STUDENT_NO_RE.match(student_no):
        errors.insert(0, '学号只能是字母或数字，最多 20 位')
    else:
        student = StudentModel.get_by_no(student_no)
        if not student:
            errors.insert(0, f'学号 {student_no} 不在学生名册中，请先到"学生名册"添加')
        else:
            record['student_name'] = student['student_name']

    if not date_raw:
        errors.append('日期不能为空')
    else:
        d = parse_date(date_raw)
        if d is None:
            errors.append('日期格式不正确，应为 YYYY-MM-DD')
        else:
            record['date_str'] = d.isoformat()

    return record, errors


def read_filters(args):
    """从查询参数里读取列表筛选条件，非法值直接忽略。"""
    start, end = parse_date(args.get('start')), parse_date(args.get('end'))
    return {
        'start': start.isoformat() if start else '',
        'end': end.isoformat() if end else '',
        'status': args.get('status') if args.get('status') in STATUSES else '',
        'keyword': (args.get('keyword') or '').strip()[:50],
    }


def _int_arg(args, key, default, lo, hi):
    try:
        return min(max(int(args.get(key, default)), lo), hi)
    except (TypeError, ValueError):
        return default


class AttendanceController:
    @staticmethod
    def _form_data():
        return request.form if request.form else (request.get_json(silent=True) or {})

    # ---------------- 列表 ----------------
    @staticmethod
    @login_required()
    def data_list():
        return render_template('back/data_list.html', statuses=STATUSES)

    @staticmethod
    @login_required(api=True)
    def get_data_list():
        filters = read_filters(request.args)
        page = _int_arg(request.args, 'page', 1, 1, 100000)
        page_size = _int_arg(request.args, 'page_size', 20, 5, 100)
        try:
            rows, total = AttendanceModel.search(filters, page, page_size)
            return jsonify({'code': 200, 'msg': 'success', 'data': rows,
                            'total': total, 'page': page, 'page_size': page_size})
        except Exception:
            # 具体错误只写日志，不把数据库报错原文返回给页面
            logger.exception('读取考勤列表失败')
            return jsonify({'code': 500, 'msg': '读取数据失败，请稍后重试'})

    # ---------------- 导出 Excel ----------------
    @staticmethod
    @login_required()
    def export_excel():
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        filters = read_filters(request.args)
        rows = AttendanceModel.export_rows(filters)

        wb = Workbook()
        ws = wb.active
        ws.title = '考勤记录'
        headers = ['日期', '学号', '姓名', '班级', '状态', '时间', '备注']
        ws.append(headers)
        for r in rows:
            ws.append([r['date'], r['student_no'], r['student_name'], r['class_name'],
                       r['status'], r['time'], r['remark']])
            # openpyxl 会把以 = 开头的字符串当公式写入；备注是用户输入，强制按文本存，防止公式注入
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cell.data_type = 's'
        head_fill = PatternFill('solid', fgColor='2C8F2C')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal='center')
        for col, width in zip('ABCDEFG', (12, 12, 10, 10, 8, 10, 30)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

        # 第二个工作表：按状态汇总，方便直接拿去汇报
        summary = wb.create_sheet('汇总')
        summary.append(['状态', '次数'])
        for s in STATUSES:
            summary.append([s, sum(1 for r in rows if r['status'] == s)])
        summary.append(['合计', len(rows)])
        for cell in summary[1]:
            cell.font = Font(bold=True)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        span = '_'.join(x for x in (filters['start'], filters['end']) if x) or date.today().isoformat()
        filename = f'考勤记录_{span}.xlsx'
        resp = send_file(buf, as_attachment=True, download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        # 中文文件名按 RFC 5987 编码，避免部分浏览器乱码
        resp.headers['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(filename)}"
        return resp

    # ---------------- 单条新增 / 编辑 / 删除 ----------------
    @staticmethod
    @login_required()
    def data_add():
        return render_template('back/data_add.html', statuses=STATUSES,
                               students=StudentModel.list_all(), today=date.today().isoformat())

    @staticmethod
    @login_required(api=True)
    def add_data():
        record, errors = validate_record(AttendanceController._form_data())
        if errors:
            return jsonify({'code': 400, 'msg': '；'.join(errors)})
        if AttendanceModel.exists_on_date(record['student_no'], record['date_str']):
            return jsonify({'code': 409, 'msg': DUPLICATE_MSG})
        try:
            AttendanceModel.add(**record)
            return jsonify({'code': 200, 'msg': '保存成功'})
        except Exception as e:
            if _is_duplicate_error(e):
                return jsonify({'code': 409, 'msg': DUPLICATE_MSG})
            logger.exception('新增考勤失败')
            return jsonify({'code': 500, 'msg': '保存失败，请稍后重试'})

    @staticmethod
    @login_required()
    def data_edit(id: int):
        record = AttendanceModel.get_by_id(id)
        if not record:
            return redirect(url_for('web.attendance_list'))
        return render_template('back/data_edit.html', record=record, statuses=STATUSES,
                               students=StudentModel.list_all())

    @staticmethod
    @login_required(api=True)
    def update_data(id: int):
        if not AttendanceModel.get_by_id(id):
            return jsonify({'code': 404, 'msg': '记录不存在或已被删除'})
        record, errors = validate_record(AttendanceController._form_data())
        if errors:
            return jsonify({'code': 400, 'msg': '；'.join(errors)})
        if AttendanceModel.exists_on_date(record['student_no'], record['date_str'], exclude_id=id):
            return jsonify({'code': 409, 'msg': DUPLICATE_MSG})
        try:
            AttendanceModel.update(record_id=id, **record)
            return jsonify({'code': 200, 'msg': '修改成功'})
        except Exception as e:
            if _is_duplicate_error(e):
                return jsonify({'code': 409, 'msg': DUPLICATE_MSG})
            logger.exception('修改考勤失败')
            return jsonify({'code': 500, 'msg': '修改失败，请稍后重试'})

    @staticmethod
    @login_required(api=True)
    def delete_data(id: int):
        try:
            AttendanceModel.delete(id)
            return jsonify({'code': 200, 'msg': '删除成功'})
        except Exception:
            logger.exception('删除考勤失败')
            return jsonify({'code': 500, 'msg': '删除失败，请稍后重试'})

    # ---------------- 批量点名 ----------------
    @staticmethod
    @login_required()
    def roll_call():
        d = parse_date(request.args.get('date')) or date.today()
        return render_template('back/roll_call.html', statuses=STATUSES, date=d.isoformat())

    @staticmethod
    @login_required(api=True)
    def get_roll_call():
        d = parse_date(request.args.get('date'))
        if d is None:
            return jsonify({'code': 400, 'msg': '日期格式不正确，应为 YYYY-MM-DD'})
        try:
            return jsonify({'code': 200, 'msg': 'success', 'data': AttendanceModel.roll_call_sheet(d.isoformat())})
        except Exception:
            logger.exception('读取点名表失败')
            return jsonify({'code': 500, 'msg': '读取数据失败，请稍后重试'})

    @staticmethod
    @login_required(api=True)
    def save_roll_call():
        data = request.get_json(silent=True) or {}
        d = parse_date(data.get('date'))
        if d is None:
            return jsonify({'code': 400, 'msg': '日期格式不正确，应为 YYYY-MM-DD'})
        items = data.get('items')
        if not isinstance(items, list) or not items:
            return jsonify({'code': 400, 'msg': '没有要保存的学生'})

        roster = {s['student_no']: s['student_name'] for s in StudentModel.list_all()}
        cleaned, errors, seen = [], [], set()
        for it in items:
            if not isinstance(it, dict):
                return jsonify({'code': 400, 'msg': '数据格式不正确'})
            no = (it.get('student_no') or '').strip()
            if no not in roster:
                errors.append(f'学号 {no or "(空)"} 不在名册中')
                continue
            if no in seen:
                errors.append(f'学号 {no} 重复提交')
                continue
            seen.add(no)
            entry, errs = validate_entry(it)
            if errs:
                errors.append(f'{roster[no]}：' + '，'.join(errs))
                continue
            cleaned.append({'student_no': no, 'student_name': roster[no], **entry})
        if errors:
            return jsonify({'code': 400, 'msg': '；'.join(errors)})
        try:
            AttendanceModel.save_roll_call(d.isoformat(), cleaned)
            return jsonify({'code': 200, 'msg': f'已保存 {d.isoformat()} 共 {len(cleaned)} 名学生的考勤'})
        except Exception:
            logger.exception('保存点名失败')
            return jsonify({'code': 500, 'msg': '保存失败，请稍后重试'})
