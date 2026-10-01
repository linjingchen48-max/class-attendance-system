"""端到端测试：真实 MySQL + 正在运行的服务。

用法（先导入 schema.sql 并启动服务）：
    python run.py            # 另开一个终端
    python e2e_test.py       # 默认连 http://127.0.0.1:5000，可用 BASE_URL 覆盖

测试数据都用 2020 年的日期和 E2E 开头的学号，结束时自动清理，不影响演示数据。
"""
import io
import os
import sys

import requests

B = os.getenv("BASE_URL", "http://127.0.0.1:5000")
s = requests.Session()
fails = 0


def check(label, cond, extra=""):
    global fails
    fails += (not cond)
    print(f"  {'✓' if cond else '✗'} {label}" + (f" — {extra}" if extra else ""))


def post(url, **data):
    return s.post(B + url, json=data).json()


def get(url, **params):
    return s.get(B + url, params=params).json()


def student_id(no):
    return next((x["id"] for x in get("/back/student/api/list")["data"] if x["student_no"] == no), None)


def record_ids(**filters):
    return [r["id"] for r in get("/back/attendance/api/list", page_size=100, **filters)["data"]]


def cleanup():
    for rid in record_ids(start="2020-01-01", end="2020-12-31"):
        post(f"/back/attendance/api/delete/{rid}")
    for no in ("E2E001", "E2E002", "E2E003", "E2E009"):
        sid = student_id(no)
        if sid:
            post(f"/back/student/api/delete/{sid}")


r = s.post(B + "/back/login", data={"username": "admin", "password": "admin123"}).json()
if r.get("code") != 200:
    sys.exit(f"登录失败：{r}")
cleanup()

print("── 学生名册（第 4 条）")
check("添加学生", post("/back/student/api/add", student_no="E2E001", student_name="测试甲", class_name="测试班")["code"] == 200)
check("添加学生 2", post("/back/student/api/add", student_no="E2E002", student_name="测试乙", class_name="测试班")["code"] == 200)
check("学号重复 → 拒绝", post("/back/student/api/add", student_no="E2E001", student_name="X")["code"] == 409)
check("学号非法 → 拒绝", post("/back/student/api/add", student_no="E2E-1", student_name="X")["code"] == 400)

print("── 单条录入（第 1、2、4 条）")
base = dict(student_no="E2E001", student_name="前端乱填的名字", date="2020-01-06", status="迟到", time="08:05", remark="堵车")
r = post("/back/attendance/api/add", **base); check("新增成功", r["code"] == 200, r["msg"])
row = get("/back/attendance/api/list", keyword="E2E001")["data"][0]
check("姓名取自名册，不信任前端提交", row["student_name"] == "测试甲", row["student_name"])
r = post("/back/attendance/api/add", **base); check("同人同天重复 → 409", r["code"] == 409, r["msg"])
r = post("/back/attendance/api/add", **{**base, "student_no": "E2E999"}); check("名册外学号 → 拒绝", r["code"] == 400, r["msg"])
r = post("/back/attendance/api/add", **{**base, "date": "2020-02-30"}); check("不存在的日期 → 拒绝", r["code"] == 400, r["msg"])
r = post("/back/attendance/api/add", **{**base, "status": "早退"}); check("非法状态 → 拒绝", r["code"] == 400, r["msg"])

print("── 名册改名 / 改学号联动（第 4 条）")
sid = student_id("E2E001")
r = post(f"/back/student/api/update/{sid}", student_no="E2E003", student_name="测试甲改", class_name="测试班")
check("改名 + 改学号", r["code"] == 200, r["msg"])
rows = get("/back/attendance/api/list", keyword="E2E003")["data"]
check("考勤记录跟着换学号、显示新名字", len(rows) == 1 and rows[0]["student_name"] == "测试甲改",
      str([(x["student_no"], x["student_name"]) for x in rows]))
r = post(f"/back/student/api/delete/{sid}"); check("有考勤记录的学生不能删", r["code"] == 409, r["msg"])
r = post(f"/back/student/api/update/{sid}", student_no="E2E002", student_name="撞号")
check("改成别人的学号 → 拒绝", r["code"] == 409, r["msg"])

print("── 批量点名（第 13 条）")
sheet = get("/back/attendance/api/roll-call", date="2020-01-06")["data"]
mine = {x["student_no"]: x for x in sheet if x["student_no"].startswith("E2E")}
check("点名表包含全部名册学生", {"E2E003", "E2E002"} <= set(mine), str(list(mine)))
check("已有记录带出原状态，没记录的为空", mine["E2E003"]["status"] == "迟到" and mine["E2E002"]["status"] is None)
items = [{"student_no": x["student_no"], "status": "正常", "time": "", "remark": ""} for x in sheet]
items = [it if it["student_no"] != "E2E002" else {**it, "status": "旷课", "remark": "<b>未到</b>"} for it in items]
r = post("/back/attendance/api/roll-call", date="2020-01-06", items=items); check("全班保存", r["code"] == 200, r["msg"])
r = post("/back/attendance/api/roll-call", date="2020-01-06", items=items); check("重复保存不产生重复记录", r["code"] == 200)
day = get("/back/attendance/api/list", start="2020-01-06", end="2020-01-06", page_size=100)
check("当天记录数 = 名册人数", day["total"] == len(sheet), f"{day['total']} vs {len(sheet)}")
check("已有记录被覆盖（迟到 → 正常）", any(x["student_no"] == "E2E003" and x["status"] == "正常" for x in day["data"]))
r = post("/back/attendance/api/roll-call", date="2020-01-06", items=[{"student_no": "E2E002", "status": "早退"}])
check("点名里的非法状态 → 拒绝", r["code"] == 400, r["msg"])

print("── 筛选、分页（第 10 条）")
r = get("/back/attendance/api/list", start="2020-01-06", end="2020-01-06", status="旷课")
check("按日期 + 状态筛选", r["total"] == 1 and r["data"][0]["student_no"] == "E2E002", f"total={r['total']}")
r = get("/back/attendance/api/list", keyword="测试乙")
check("按姓名关键字", r["total"] >= 1 and all("测试乙" in x["student_name"] for x in r["data"]))
r = get("/back/attendance/api/list", keyword="%")
check("关键字 % 不当通配符", r["total"] == 0, f"total={r['total']}")
p1 = get("/back/attendance/api/list", start="2020-01-06", end="2020-01-06", page_size=5, page=1)
p2 = get("/back/attendance/api/list", start="2020-01-06", end="2020-01-06", page_size=5, page=2)
check("分页不重不漏", len({x["id"] for x in p1["data"]} & {x["id"] for x in p2["data"]}) == 0
      and len(p1["data"]) + len(p2["data"]) == min(p1["total"], 10))

print("── 导出 Excel（第 14 条）")
resp = s.get(B + "/back/attendance/export", params={"start": "2020-01-06", "end": "2020-01-06"})
check("返回 xlsx 文件", resp.status_code == 200 and resp.content[:2] == b"PK", resp.headers.get("Content-Type", ""))
check("中文文件名", "filename*=UTF-8''" in resp.headers.get("Content-Disposition", ""))
try:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb["考勤记录"]
    check("行数与筛选结果一致", ws.max_row - 1 == day["total"], f"{ws.max_row - 1} 行")
    check("含汇总表", "汇总" in wb.sheetnames)
except ImportError:
    print("  · 未安装 openpyxl，跳过内容校验")

print("── 个人统计（第 15 条）")
html = s.get(B + "/back/student/E2E002", params={"month": "2020-01"}).text
check("个人页显示旷课 1 次", "旷课" in html and "1 次" in html)
check("个人页转义备注", "&lt;b&gt;未到&lt;/b&gt;" in html and "<b>未到</b>" not in html)

print("── 大屏未登记（第 3 条）")
st = requests.get(B + "/front/api/stats").json()["data"]
ta = st["today_attendance"]
recorded = ta["present"] + ta["late"] + ta["leave"] + ta["absent"]
check("应到 = 已登记 + 未登记", st["total_students"] == recorded + ta["unregistered"],
      f"{st['total_students']} = {recorded} + {ta['unregistered']}")
check("未登记的学生出现在异常名单", sum(1 for x in st["abnormal_list"] if x["type"] == "未登记") == ta["unregistered"])

print("── 未登录")
anon = requests.Session()
check("接口返回 401", anon.get(B + "/back/student/api/list").json()["code"] == 401)
check("导出跳登录页", anon.get(B + "/back/attendance/export", allow_redirects=False).status_code == 302)

cleanup()
check("测试数据已清理", not record_ids(start="2020-01-01", end="2020-12-31") and not student_id("E2E002"))
print(f"\n{'全部通过' if not fails else str(fails) + ' 项失败'}")
sys.exit(1 if fails else 0)
